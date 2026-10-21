import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


SCRIPT = Path(__file__).with_name('convert_to_md.py')


class ConversionTests(unittest.TestCase):
    def run_cli(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)],
                              capture_output=True, text=True, encoding='utf-8')

    def test_docx_order_table_escaping_and_no_overwrite(self):
        with tempfile.TemporaryDirectory(dir=SCRIPT.parent) as temp:
            root = Path(temp)
            source = root / '中文.docx'
            doc = Document()
            doc.add_heading('中文标题', 1)
            doc.add_paragraph('先看正文')
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = '项目'
            table.cell(0, 1).text = '说明'
            table.cell(1, 0).text = 'A|B'
            table.cell(1, 1).text = '第一行\n第二行'
            doc.add_paragraph('结束', style='List Bullet')
            doc.save(source)
            output = root / 'output'
            result = self.run_cli(source, '-o', output)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            target = output / '中文.docx.md'
            md = target.read_text(encoding='utf-8')
            self.assertIn('# 中文标题', md)
            self.assertIn('A\\|B', md)
            self.assertIn('第一行<br>第二行', md)
            self.assertLess(md.index('先看正文'), md.index('| 项目'))
            self.assertIn('- 结束', md)
            target.write_text('保留已有结果', encoding='utf-8')
            self.assertEqual(self.run_cli(source, '-o', output).returncode, 0)
            self.assertEqual(target.read_text(encoding='utf-8'), '保留已有结果')

    def test_pdf_example_has_chinese_text(self):
        source = SCRIPT.parent / 'ziliao/02_PDF合作案例/01_离散制造工厂设备联网.pdf'
        if not source.exists():
            self.skipTest('example file unavailable')
        with tempfile.TemporaryDirectory(dir=SCRIPT.parent) as temp:
            result = self.run_cli(source, '-o', temp)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            md = (Path(temp) / (source.name + '.md')).read_text(encoding='utf-8')
            self.assertIn('设备', md)
            self.assertIn('<!-- 第 1 页 -->', md)

    def test_batch_continues_after_corrupt_file_and_preserves_subfolders(self):
        with tempfile.TemporaryDirectory(dir=SCRIPT.parent) as temp:
            root = Path(temp)
            folder = root / 'input' / '子目录'
            folder.mkdir(parents=True)
            doc = Document()
            doc.add_paragraph('正常文件')
            doc.save(folder / 'good.docx')
            (folder / 'bad.pdf').write_bytes(b'invalid pdf')
            result = self.run_cli(root / 'input', '-r', '-o', root / 'out')
            self.assertEqual(result.returncode, 1)
            self.assertTrue((root / 'out/子目录/good.docx.md').exists())
            self.assertFalse((root / 'out/子目录/bad.pdf.md').exists())


if __name__ == '__main__':
    unittest.main()
