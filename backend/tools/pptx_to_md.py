"""PPTX 各页标题、文字、表格与图片占位转 Markdown。"""
from backend.tools.common import run_cli, check_archive, table


def convert(path):
    check_archive(path)
    from pptx import Presentation
    book = Presentation(path)
    if len(book.slides) > 1000:
        raise ValueError('幻灯片页数过多')
    parts = []
    found = False
    for index, slide in enumerate(book.slides, 1):
        body = []
        for shape in slide.shapes:
            if shape.has_table:
                body.append(table([[cell.text for cell in row.cells] for row in shape.table.rows]))
            elif shape.has_text_frame and shape.text.strip():
                body.append(shape.text.strip())
            elif shape.shape_type == 13:
                body.append('[图片占位]')
        if body:
            found = True
        parts.append(f'# 第 {index} 页\n\n' + '\n\n'.join(body))
    if not found:
        raise ValueError('幻灯片无可提取内容')
    return '\n\n'.join(parts)


def main():
    run_cli(convert)


if __name__ == '__main__':
    main()
