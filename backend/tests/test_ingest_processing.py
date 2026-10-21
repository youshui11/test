"""真实转换器与 Markdown 语义切分行为。"""
from pathlib import Path
import pytest


def test_semantic_chunks_keep_headings_overlap_and_limits():
    from backend.services.chunker import chunk_markdown
    paragraph = '这是一句用于验证语义切分和资料上下文的文字。' * 14
    markdown = '# 概述\n\n## 背景\n\n' + '\n\n'.join([paragraph] * 4) + '\n\n## 目标\n\n这是目标。'
    chunks = chunk_markdown(markdown, '说明.md')
    assert len(chunks) >= 3
    assert chunks[0].metadata['heading_path'] == '概述 > 背景'
    assert all(0 < len(c.text) <= 800 and c.metadata['char_count'] == len(c.text) for c in chunks)
    assert [c.metadata['chunk_index'] for c in chunks] == list(range(len(chunks)))
    assert all(c.metadata['source_file'] == '说明.md' for c in chunks)
    assert chunks[-1].metadata['heading_path'] == '概述 > 目标'
    assert '这是目标' in chunks[-1].text and '用于验证' not in chunks[-1].text
    assert chunks[0].text.endswith('这是一句用于验证语义切分和资料上下文的文字。')
    assert '这是一句用于验证语义切分和资料上下文的文字。' in chunks[1].text


def test_long_paragraph_table_and_code_are_preserved():
    from backend.services.chunker import chunk_markdown
    long = chunk_markdown('# 无标点\n' + '中' * 2500, 'long.md')
    assert len(long) >= 4 and all(len(c.text) <= 800 for c in long)
    assert sum(c.text.count('中') for c in long) >= 2500
    markdown = '# 数据\n\n| 字段 | 值 |\n| --- | --- |\n| 甲 | 乙 |\n\n```python\n# 不是标题\nprint(1)\n```'
    result = chunk_markdown(markdown, 'data.md')
    assert len(result) == 1
    assert result[0].metadata['heading_path'] == '数据'
    assert '| 甲 | 乙 |' in result[0].text and '```python' in result[0].text
    assert chunk_markdown(' \n\x00 ', 'empty.md') == []


@pytest.mark.asyncio
async def test_text_conversion_and_invalid_file(tmp_path):
    from backend.services.converter import convert_file, ConvertError
    source = tmp_path / '中文.txt'
    source.write_text('中文资料\n第二行', encoding='utf-8')
    assert '中文资料' in await convert_file(source)
    with pytest.raises(ConvertError, match='不支持的文件格式'):
        await convert_file(tmp_path / 'bad.exe')
    source.write_bytes(b'')
    with pytest.raises(ConvertError):
        await convert_file(source)


@pytest.mark.asyncio
async def test_conversion_drains_large_output_without_deadlock(tmp_path):
    from backend.services.converter import convert_file
    source = tmp_path / 'large.md'
    source.write_text('large document text\n' * 15000, encoding='utf-8')
    text = await convert_file(source, timeout=3)
    assert text.count('large document text') == 15000


def test_distinct_sentences_overlap_without_losing_content():
    from backend.services.chunker import chunk_markdown, sentences
    source = [f'第{i}条资料说明了不同的业务知识和处理步骤。' for i in range(100)]
    chunks = chunk_markdown('# 测试\n\n' + ''.join(source), 'unique.md')
    assert len(chunks) > 2
    for previous, current in zip(chunks, chunks[1:]):
        assert sentences(previous.text.split('\n\n', 1)[1])[-1] in current.text
    assert all(any(sentence in chunk.text for chunk in chunks) for sentence in source)


def test_nested_fences_do_not_promote_code_to_headings():
    from backend.services.chunker import chunk_markdown
    markdown = '# C#\n\n````markdown\n```python\n# keep-as-code\nprint(1)\n```\n````\n\n## 正文\n\n说明。'
    chunks = chunk_markdown(markdown, 'fences.md')
    assert [c.metadata['heading_path'] for c in chunks] == ['C#', 'C# > 正文']
    assert '# keep-as-code' in chunks[0].text
    assert '````markdown' in chunks[0].text and chunks[0].text.endswith('````')


@pytest.mark.asyncio
async def test_office_and_pdf_conversion(tmp_path):
    from backend.services.converter import convert_file, ConvertError
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    import pymupdf
    doc = Document()
    doc.add_heading('文档标题', 1)
    doc.add_paragraph('文档正文')
    doc.save(tmp_path / 'test.docx')
    assert '# 文档标题' in await convert_file(tmp_path / 'test.docx')
    wb = Workbook()
    wb.active.title = '销售'
    wb.active.append(['项目', '金额'])
    wb.active.append(['A|B', 12])
    wb.save(tmp_path / 'test.xlsx')
    excel = await convert_file(tmp_path / 'test.xlsx')
    assert '销售' in excel and 'A\\|B' in excel and '12' in excel
    import xlwt
    legacy = xlwt.Workbook()
    sheet = legacy.add_sheet('旧版资料')
    sheet.write(0, 0, '名称')
    sheet.write(1, 0, '旧版数据')
    legacy.save(str(tmp_path / 'test.xls'))
    assert '旧版数据' in await convert_file(tmp_path / 'test.xls')
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[1])
    slide.shapes.title.text = '幻灯片标题'
    slide.placeholders[1].text = '幻灯片正文'
    ppt.save(tmp_path / 'test.pptx')
    assert '幻灯片正文' in await convert_file(tmp_path / 'test.pptx')
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), 'PDF sample content for ingestion')
        pdf.save(tmp_path / 'test.pdf')
    assert 'PDF sample content' in await convert_file(tmp_path / 'test.pdf')
    (tmp_path / 'broken.docx').write_bytes(b'invalid')
    with pytest.raises(ConvertError):
        await convert_file(tmp_path / 'broken.docx')
