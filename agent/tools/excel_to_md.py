"""XLSX/XLS 每个 sheet 转为 Markdown 表格，不执行公式或宏。"""
from agent.tools.common import run_cli, check_archive, table, MAX_TEXT_CHARS


def convert(path):
    parts = []
    if path.suffix.lower() == '.xlsx':
        check_archive(path)
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row > 100_000 or sheet.max_column > 256:
                    raise ValueError('工作表尺寸过大')
                parts.append('# ' + sheet.title + '\n\n' + table(sheet.iter_rows(values_only=True)))
        finally:
            book.close()
    else:
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows > 100_000 or sheet.ncols > 256:
                    raise ValueError('工作表尺寸过大')
                parts.append('# ' + sheet.name + '\n\n' + table(sheet.row_values(i) for i in range(sheet.nrows)))
        finally:
            book.release_resources()
    result = '\n\n'.join(parts)
    if len(result) > MAX_TEXT_CHARS:
        raise ValueError('工作簿文本过长')
    if not any('|' in part for part in parts):
        raise ValueError('工作簿没有数据')
    return result


def main():
    run_cli(convert)


if __name__ == '__main__':
    main()
